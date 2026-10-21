import pptx

prs = pptx.Presentation(r'C:/Users/Shivam Keshari/.gemini/antigravity/brain/8ef3ef34-7b4d-453e-b498-75b3644d3aa7/.user_uploaded/media_1792515573274_65da0e82.pptx')
for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for j, shape in enumerate(slide.shapes):
        shape_type = str(shape.shape_type)
        text = shape.text.strip().replace('\n', ' // ') if shape.has_text_frame else ''
        has_table = shape.has_table
        table_info = f'Table {len(shape.table.rows)}x{len(shape.table.columns)}' if has_table else ''
        print(f'  Shape {j}: type={shape_type}, text="{text[:80]}", {table_info}')
        if shape.has_table:
            for r_idx, row in enumerate(shape.table.rows):
                row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(f'    Row {r_idx}: {row_vals}')
